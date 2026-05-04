import io
import os

from dotenv import load_dotenv
load_dotenv()

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse
from openai import OpenAI
import re
from datetime import date

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pypdf import PdfReader
from pydantic import BaseModel


app = FastAPI(title="Aman.ai API")

app.add_middleware(
	CORSMiddleware,
	allow_origins=["*"],
	allow_credentials=False,
	allow_methods=["*"],
	allow_headers=["*"],
)

OPENAI_MODEL = os.getenv("OPENAI_MODEL", "gpt-4o")

PHASE_INSTRUCTIONS = {
	"Requirements": (
		"You are Aman.ai, an AI-Driven Software Security Analyst Assistant for Secure SDLC.\n\n"
		"The user has uploaded a Requirements phase document (BRD, SRS, Functional Requirements, etc.).\n\n"
		"You must:\n"
		"- Analyze the requirements document\n"
		"- Generate security requirements\n"
		"- Identify missing security controls\n"
		"- Map findings to OWASP Top 10\n"
		"- Highlight authentication, authorization, validation, and logging needs\n\n"
		"Output a professional Requirements Phase Security Report with these sections:\n"
		"1. Short Summary (5-7 lines, simple and clear for non-technical users)\n"
		"2. Security Requirements\n"
		"3. Missing Controls\n"
		"4. OWASP Top 10 Mapping\n"
		"5. Recommendations\n\n"
		"If the uploaded document is too vague or generic with no real content, respond with:\n"
		"'Input too generic. Please upload the required document for the selected SDLC phase.'\n\n"
		"Do NOT return long paragraphs. Be structured, professional, and formal."
	),
	"Design": (
		"You are Aman.ai, an AI-Driven Software Security Analyst Assistant for Secure SDLC.\n\n"
		"The user has uploaded a Design phase document (UML, Use Case Diagrams, DFD, Architecture Diagram, Database Design, Network Design).\n\n"
		"You must:\n"
		"- Perform Threat Modeling\n"
		"- Apply STRIDE methodology\n"
		"- Identify trust boundaries\n"
		"- Suggest secure architecture improvements\n"
		"- Recommend controls using OWASP ASVS\n"
		"- Review secure design principles\n\n"
		"Output a professional Design Phase Security Report with these sections:\n"
		"1. Short Summary (5-7 lines, simple and clear for non-technical users)\n"
		"2. Threat Table\n"
		"3. STRIDE Analysis\n"
		"4. Security Controls\n"
		"5. Secure Architecture Recommendations\n\n"
		"If the uploaded document is too vague or generic with no real content, respond with:\n"
		"'Input too generic. Please upload the required document for the selected SDLC phase.'\n\n"
		"Do NOT return long paragraphs. Be structured, professional, and formal."
	),
	"Development": (
		"You are Aman.ai, an AI-Driven Software Security Analyst Assistant for Secure SDLC.\n\n"
		"The user has uploaded a Development phase document (Source Code, API Design, Backend Logic, Database Schema, Config Files, Authentication Flow, Frontend Integration).\n\n"
		"You must:\n"
		"- Review secure coding practices\n"
		"- Identify code security weaknesses\n"
		"- Review API security\n"
		"- Review authentication and authorization\n"
		"- Review FastAPI backend security if applicable\n"
		"- Review React frontend security if applicable\n"
		"- Review SQL schema security\n"
		"- Validate input handling\n"
		"- Highlight missing protections\n\n"
		"Output a professional Development Phase Security Report with these sections:\n"
		"1. Short Summary (5-7 lines, simple and clear for non-technical users)\n"
		"2. Secure Coding Issues\n"
		"3. API Security Review\n"
		"4. Authentication Review\n"
		"5. Validation Issues\n"
		"6. Database Security Review\n"
		"7. Recommended Fixes\n\n"
		"If the uploaded document is too vague or generic with no real content, respond with:\n"
		"'Input too generic. Please upload the required document for the selected SDLC phase.'\n\n"
		"Do NOT return long paragraphs. Be structured, professional, and formal."
	),
	"Testing": (
		"You are Aman.ai, an AI-Driven Software Security Analyst Assistant for Secure SDLC.\n\n"
		"The user has uploaded a Testing phase document (Test Cases, QA Reports, Security Testing Reports, UAT Reports, Penetration Testing Results, Validation Reports).\n\n"
		"You must:\n"
		"- Review testing completeness\n"
		"- Identify missing security tests\n"
		"- Suggest abuse cases\n"
		"- Suggest negative testing scenarios\n"
		"- Validate OWASP Testing coverage\n"
		"- Review authentication and access control testing\n\n"
		"Output a professional Testing Phase Security Report with these sections:\n"
		"1. Short Summary (5-7 lines, simple and clear for non-technical users)\n"
		"2. Test Case Review\n"
		"3. Missing Security Tests\n"
		"4. Abuse Cases\n"
		"5. Negative Testing\n"
		"6. Security Testing Recommendations\n\n"
		"If the uploaded document is too vague or generic with no real content, respond with:\n"
		"'Input too generic. Please upload the required document for the selected SDLC phase.'\n\n"
		"Do NOT return long paragraphs. Be structured, professional, and formal."
	),
	"Deployment": (
		"You are Aman.ai, an AI-Driven Software Security Analyst Assistant for Secure SDLC.\n\n"
		"The user has uploaded a Deployment phase document (Deployment Config, Server Config, CI/CD Config, Network Config, Cloud Settings, Security Policies).\n\n"
		"You must:\n"
		"- Review deployment security\n"
		"- Check HTTPS configuration\n"
		"- Review secure headers\n"
		"- Review logging and monitoring setup\n"
		"- Recommend IDS/IPS\n"
		"- Suggest alerting mechanisms\n"
		"- Review API exposure\n"
		"- Recommend rate limiting\n\n"
		"Output a professional Deployment Phase Security Report with these sections:\n"
		"1. Short Summary (5-7 lines, simple and clear for non-technical users)\n"
		"2. Deployment Security Review\n"
		"3. Monitoring Recommendations\n"
		"4. Logging Requirements\n"
		"5. Incident Alerting\n"
		"6. Secure Deployment Recommendations\n\n"
		"If the uploaded document is too vague or generic with no real content, respond with:\n"
		"'Input too generic. Please upload the required document for the selected SDLC phase.'\n\n"
		"Do NOT return long paragraphs. Be structured, professional, and formal."
	),
	"Maintenance": (
		"You are Aman.ai, an AI-Driven Software Security Analyst Assistant for Secure SDLC.\n\n"
		"The user has uploaded a Maintenance phase document (Patch Reports, Vulnerability Reports, Monitoring Logs, Audit Reports, Incident Reports, Dependency Review Reports).\n\n"
		"You must:\n"
		"- Review security updates\n"
		"- Review vulnerability management\n"
		"- Suggest dependency updates\n"
		"- Recommend continuous improvement\n"
		"- Review audit logs\n"
		"- Suggest backup and recovery improvements\n"
		"- Recommend periodic security reviews\n\n"
		"Output a professional Maintenance Phase Security Report with these sections:\n"
		"1. Short Summary (5-7 lines, simple and clear for non-technical users)\n"
		"2. Update Recommendations\n"
		"3. Vulnerability Review\n"
		"4. Dependency Security Review\n"
		"5. Monitoring Improvements\n"
		"6. Backup & Recovery Recommendations\n\n"
		"If the uploaded document is too vague or generic with no real content, respond with:\n"
		"'Input too generic. Please upload the required document for the selected SDLC phase.'\n\n"
		"Do NOT return long paragraphs. Be structured, professional, and formal."
	),
}


def _extract_text(raw_bytes: bytes, filename: str, content_type: str | None) -> str:
	fname = (filename or "").lower()

	if fname.endswith(".pdf") or content_type == "application/pdf":
		try:
			reader = PdfReader(io.BytesIO(raw_bytes))
			text = "\n".join(page.extract_text() or "" for page in reader.pages)
			return text.strip() or "[PDF contained no extractable text.]"
		except Exception:
			return "[Failed to extract PDF text.]"

	if fname.endswith(".docx"):
		try:
			doc = Document(io.BytesIO(raw_bytes))
			text = "\n".join(p.text for p in doc.paragraphs)
			return text.strip() or "[Word document contained no extractable text.]"
		except Exception:
			return "[Failed to extract Word document text.]"

	# Plain text fallback
	try:
		return raw_bytes.decode("utf-8").strip()
	except UnicodeDecodeError:
		return raw_bytes.decode("latin-1", errors="replace").strip()


def _get_client() -> OpenAI:
	api_key = os.getenv("OPENAI_API_KEY")
	if not api_key:
		raise HTTPException(status_code=500, detail="Missing OPENAI_API_KEY environment variable.")
	return OpenAI(api_key=api_key)


# ── Colour palette ────────────────────────────────────────────────────────────
_DARK_BLUE  = RGBColor(0x1F, 0x39, 0x7A)
_MID_BLUE   = RGBColor(0x2E, 0x4A, 0x8F)
_ACCENT     = RGBColor(0x3B, 0x82, 0xF6)
_WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
_LIGHT_GREY = RGBColor(0xF2, 0xF2, 0xF2)
_DARK_GREY  = RGBColor(0x40, 0x40, 0x40)


def _shade_cell(cell, hex_color: str):
	"""Fill a table cell with a solid background colour."""
	tc = cell._tc
	tcPr = tc.get_or_add_tcPr()
	shd = OxmlElement("w:shd")
	shd.set(qn("w:val"), "clear")
	shd.set(qn("w:color"), "auto")
	shd.set(qn("w:fill"), hex_color)
	tcPr.append(shd)


def _set_cell_border(cell, **kwargs):
	"""Add borders to a table cell. kwargs: top, bottom, left, right."""
	tc = cell._tc
	tcPr = tc.get_or_add_tcPr()
	tcBorders = OxmlElement("w:tcBorders")
	for edge, color in kwargs.items():
		border = OxmlElement(f"w:{edge}")
		border.set(qn("w:val"), "single")
		border.set(qn("w:sz"), "4")
		border.set(qn("w:space"), "0")
		border.set(qn("w:color"), color)
		tcBorders.append(border)
	tcPr.append(tcBorders)


def _add_inline_formatted_run(paragraph, text: str, base_size: int = 11):
	"""Render inline **bold**, *italic*, ***bold-italic***, and `code` spans."""
	pattern = re.compile(
		r"(\*\*\*[^*]+\*\*\*"		# ***bold-italic***
		r"|\*\*[^*]+\*\*"			# **bold**
		r"|\*[^*]+\*"				# *italic*
		r"|`[^`]+`)"					# `code`
	)
	for part in pattern.split(text):
		if not part:
			continue
		run = paragraph.add_run()
		run.font.size = Pt(base_size)
		run.font.color.rgb = _DARK_GREY
		if part.startswith("***") and part.endswith("***"):
			run.text = part[3:-3]
			run.bold = True
			run.italic = True
		elif part.startswith("**") and part.endswith("**"):
			run.text = part[2:-2]
			run.bold = True
		elif part.startswith("*") and part.endswith("*"):
			run.text = part[1:-1]
			run.italic = True
		elif part.startswith("`") and part.endswith("`"):
			run.text = part[1:-1]
			run.font.name = "Courier New"
			run.font.size = Pt(base_size - 1)
			run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
		else:
			run.text = part


def _is_table_row(line: str) -> bool:
	return line.strip().startswith("|") and line.strip().endswith("|")


def _is_separator_row(line: str) -> bool:
	return bool(re.match(r"^[\|\s\-:]+$", line.strip()))


def _collect_table_block(lines: list[str], start: int):
	"""Return (rows_list, next_index). rows_list is list of list-of-cell-strings."""
	rows = []
	i = start
	while i < len(lines) and _is_table_row(lines[i]):
		if not _is_separator_row(lines[i]):
			cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
			rows.append(cells)
		i += 1
	return rows, i


def _render_table(doc: Document, rows: list[list[str]]):
	if not rows:
		return
	col_count = max(len(r) for r in rows)
	table = doc.add_table(rows=len(rows), cols=col_count)
	table.style = "Table Grid"

	for r_idx, row in enumerate(rows):
		for c_idx in range(col_count):
			cell = table.cell(r_idx, c_idx)
			text = row[c_idx] if c_idx < len(row) else ""
			cell.text = ""
			p = cell.paragraphs[0]
			p.paragraph_format.space_before = Pt(3)
			p.paragraph_format.space_after = Pt(3)
			if r_idx == 0:
				# Header row
				_shade_cell(cell, "1F397A")
				_set_cell_border(cell, top="FFFFFF", bottom="FFFFFF", left="FFFFFF", right="FFFFFF")
				run = p.add_run(text)
				run.bold = True
				run.font.size = Pt(10)
				run.font.color.rgb = _WHITE
			else:
				# Alternate row shading
				_shade_cell(cell, "F2F2F2" if r_idx % 2 == 0 else "FFFFFF")
				_set_cell_border(cell, top="CCCCCC", bottom="CCCCCC", left="CCCCCC", right="CCCCCC")
				_add_inline_formatted_run(p, text, base_size=10)
	doc.add_paragraph()


def _build_docx(title: str, content: str) -> bytes:
	doc = Document()

	# ── Page setup ────────────────────────────────────────────────────────────
	for section in doc.sections:
		section.top_margin    = Inches(1)
		section.bottom_margin = Inches(1)
		section.left_margin   = Inches(1.2)
		section.right_margin  = Inches(1.2)

	# ── Cover block ───────────────────────────────────────────────────────────
	for _ in range(3):
		doc.add_paragraph()

	title_para = doc.add_paragraph()
	title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
	run = title_para.add_run(title)
	run.bold = True
	run.font.size = Pt(22)
	run.font.color.rgb = _DARK_BLUE

	subtitle_para = doc.add_paragraph()
	subtitle_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
	sub = subtitle_para.add_run("AI-Driven Software Security Analyst Assistant")
	sub.font.size = Pt(12)
	sub.italic = True
	sub.font.color.rgb = _MID_BLUE

	date_para = doc.add_paragraph()
	date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
	dr = date_para.add_run(f"Generated: {date.today().strftime('%B %d, %Y')}")
	dr.font.size = Pt(10)
	dr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

	for _ in range(2):
		doc.add_paragraph()

	# Divider line via a 1-row, 1-col table with bottom border
	divider = doc.add_table(rows=1, cols=1)
	divider.style = "Table Grid"
	cell = divider.cell(0, 0)
	cell.text = ""
	_shade_cell(cell, "1F397A")
	_set_cell_border(cell, top="1F397A", bottom="1F397A", left="1F397A", right="1F397A")

	doc.add_page_break()

	# ── Body ──────────────────────────────────────────────────────────────────
	lines = content.split("\n")
	i = 0
	while i < len(lines):
		line = lines[i]
		stripped = line.strip()

		# Table block
		if _is_table_row(stripped):
			rows, i = _collect_table_block(lines, i)
			_render_table(doc, rows)
			continue

		if not stripped:
			doc.add_paragraph()
			i += 1
			continue

		# Horizontal rule
		if re.match(r"^[-*_]{3,}$", stripped):
			p = doc.add_paragraph()
			p.paragraph_format.space_before = Pt(4)
			p.paragraph_format.space_after = Pt(4)
			pPr = p._p.get_or_add_pPr()
			pBdr = OxmlElement("w:pBdr")
			bottom = OxmlElement("w:bottom")
			bottom.set(qn("w:val"), "single")
			bottom.set(qn("w:sz"), "6")
			bottom.set(qn("w:space"), "1")
			bottom.set(qn("w:color"), "1F397A")
			pBdr.append(bottom)
			pPr.append(pBdr)
			i += 1
			continue

		# H1
		if re.match(r"^# [^#]", stripped):
			p = doc.add_paragraph()
			p.paragraph_format.space_before = Pt(14)
			p.paragraph_format.space_after = Pt(4)
			r = p.add_run(stripped[2:].strip())
			r.bold = True
			r.font.size = Pt(18)
			r.font.color.rgb = _DARK_BLUE
			i += 1
			continue

		# H2
		if re.match(r"^## [^#]", stripped):
			p = doc.add_paragraph()
			p.paragraph_format.space_before = Pt(12)
			p.paragraph_format.space_after = Pt(3)
			r = p.add_run(stripped[3:].strip())
			r.bold = True
			r.font.size = Pt(14)
			r.font.color.rgb = _DARK_BLUE
			i += 1
			continue

		# H3
		if re.match(r"^### ", stripped):
			p = doc.add_paragraph()
			p.paragraph_format.space_before = Pt(10)
			p.paragraph_format.space_after = Pt(2)
			r = p.add_run(stripped[4:].strip())
			r.bold = True
			r.font.size = Pt(12)
			r.font.color.rgb = _MID_BLUE
			i += 1
			continue

		# H4
		if re.match(r"^#### ", stripped):
			p = doc.add_paragraph()
			p.paragraph_format.space_before = Pt(8)
			p.paragraph_format.space_after = Pt(2)
			r = p.add_run(stripped[5:].strip())
			r.bold = True
			r.italic = True
			r.font.size = Pt(11)
			r.font.color.rgb = _ACCENT
			i += 1
			continue

		# Nested bullet (two spaces or tab indent)
		if re.match(r"^(  |\t)[\-\*] ", line):
			p = doc.add_paragraph(style="List Bullet 2")
			p.paragraph_format.space_after = Pt(1)
			_add_inline_formatted_run(p, re.sub(r"^(  |\t)[\-\*] ", "", line).strip())
			i += 1
			continue

		# Bullet
		if re.match(r"^[\-\*] ", stripped):
			p = doc.add_paragraph(style="List Bullet")
			p.paragraph_format.space_after = Pt(2)
			_add_inline_formatted_run(p, stripped[2:].strip())
			i += 1
			continue

		# Numbered list
		if re.match(r"^\d+[.)]\ ", stripped):
			p = doc.add_paragraph(style="List Number")
			p.paragraph_format.space_after = Pt(2)
			text = re.sub(r"^\d+[.)]\ ", "", stripped)
			_add_inline_formatted_run(p, text.strip())
			i += 1
			continue

		# Blockquote
		if stripped.startswith("> "):
			p = doc.add_paragraph()
			p.paragraph_format.left_indent = Inches(0.4)
			p.paragraph_format.space_before = Pt(4)
			p.paragraph_format.space_after = Pt(4)
			r = p.add_run(stripped[2:].strip())
			r.italic = True
			r.font.size = Pt(11)
			r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
			i += 1
			continue

		# Normal paragraph
		p = doc.add_paragraph()
		p.paragraph_format.space_after = Pt(4)
		_add_inline_formatted_run(p, stripped)
		i += 1

	buf = io.BytesIO()
	doc.save(buf)
	buf.seek(0)
	return buf.read()


@app.post("/api/analyze")
async def analyze_sdlc_file(
	phase: str = Form(...),
	files: list[UploadFile] = File(...),
):
	if phase not in PHASE_INSTRUCTIONS:
		raise HTTPException(status_code=400, detail=f"Unknown phase: {phase}")

	if not files:
		raise HTTPException(status_code=400, detail="No files uploaded.")

	# Extract and combine text from all uploaded files
	combined_text = ""
	for file in files:
		raw_bytes = await file.read()
		if not raw_bytes:
			continue
		text = _extract_text(raw_bytes, file.filename, file.content_type)
		combined_text += f"\n\n--- File: {file.filename} ---\n{text}"

	if not combined_text.strip():
		raise HTTPException(status_code=400, detail="All uploaded files are empty.")

	client = _get_client()

	response = client.chat.completions.create(
		model=OPENAI_MODEL,
		messages=[
			{
				"role": "system",
				"content": PHASE_INSTRUCTIONS[phase],
			},
			{
				"role": "user",
				"content": (
					f"Please analyze the following {len(files)} document(s) for the {phase} phase.\n"
					f"{combined_text}"
				),
			},
		],
	)

	output_text = response.choices[0].message.content
	if not output_text:
		raise HTTPException(status_code=500, detail="No response received from model.")

	return JSONResponse({"phase": phase, "analysis": output_text.strip()})


class PhaseReport(BaseModel):
	phase: str
	analysis: str


@app.post("/api/final-report")
async def generate_final_report(phases: list[PhaseReport]):
	if not phases:
		raise HTTPException(status_code=400, detail="No phase reports provided.")

	# Build a cover section then append every phase report in full — no GPT summarization
	lines = ["# Final Comprehensive Security Report", "", "AI-Driven Software Security Analyst Assistant — Aman.ai", ""]

	for p in phases:
		lines.append(f"## {p.phase} Phase Security Report")
		lines.append("")
		lines.append(p.analysis)
		lines.append("")
		lines.append("---")
		lines.append("")

	full_report = "\n".join(lines)
	return JSONResponse({"analysis": full_report})


@app.post("/api/download-report")
async def download_report(
	phase: str = Form(...),
	content: str = Form(...),
):
	title = f"Aman.ai — {phase} Security Report"
	if phase.lower() == "final":
		title = "Aman.ai — Final Comprehensive Security Report"

	docx_bytes = _build_docx(title, content)
	filename = f"Aman_ai_{phase.replace(' ', '_')}_Report.docx"

	return StreamingResponse(
		io.BytesIO(docx_bytes),
		media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
		headers={"Content-Disposition": f"attachment; filename={filename}"},
	)
